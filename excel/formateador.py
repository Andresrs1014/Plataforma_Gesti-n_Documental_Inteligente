from pathlib import Path
import json
from openpyxl import load_workbook
from config import Config
from utils.logger import configurar_logger

logger = configurar_logger("Formateador")

def cargar_mapeo(archivo_mapeo: str) -> dict:
    """
    Carga el archivo JSON de mapeo de celdas
    
    Args:
        archivo_mapeo: Nombre del archivo de mapeo
        
    Returns:
        Diccionario con el mapeo campo -> celda
    """
    ruta_mapeo = Config.MAPPING_DIR / archivo_mapeo
    
    if not ruta_mapeo.exists():
        logger.warning(f"Archivo de mapeo no encontrado: {ruta_mapeo}")
        return {}
    
    try:
        with open(ruta_mapeo, 'r', encoding='utf-8') as f:
            mapeo = json.load(f)
        logger.info(f"Mapeo cargado: {len(mapeo)} campos")
        return mapeo
    except Exception as e:
        logger.error(f"Error al cargar mapeo: {e}")
        return {}

def llenar_procedimiento(ruta: Path, datos: dict):
    """
    Rellena el formato de PROCEDIMIENTO
    
    Args:
        ruta: Path del archivo Excel a rellenar
        datos: Diccionario con los datos
    """
    logger.info(f"Rellenando procedimiento: {ruta.name}")
    
    try:
        # Cargar mapeo
        mapeo = cargar_mapeo(Config.MAPEO_PROCEDIMIENTO)
        
        # Abrir workbook
        wb = load_workbook(ruta)
        ws = wb.active
        if ws is None:
            logger.error("No se pudo abrir la hoja activa del workbook")
            raise ValueError("Hoja activa no encontrada")
        
        # Rellenar celdas básicas según mapeo
        campos_rellenados = 0
        for campo, celda in mapeo.items():
            if campo in datos and datos[campo]:
                try:
                    ws[celda] = datos[campo]
                    campos_rellenados += 1
                except Exception as e:
                    logger.warning(f"No se pudo rellenar celda {celda}: {e}")
        
        logger.info(f"Campos rellenados: {campos_rellenados}/{len(mapeo)}")
        
        # Guardar y cerrar
        wb.save(ruta)
        wb.close()
        logger.info(f"✅ Procedimiento guardado exitosamente")
        
    except Exception as e:
        logger.error(f"Error al rellenar procedimiento: {e}")
        raise

def llenar_instructivo(ruta: Path, datos: dict):
    """
    Rellena el formato de INSTRUCTIVO
    
    Args:
        ruta: Path del archivo Excel a rellenar
        datos: Diccionario con los datos
    """
    logger.info(f"Rellenando instructivo: {ruta.name}")
    
    try:
        # Cargar mapeo
        mapeo = cargar_mapeo(Config.MAPEO_INSTRUCTIVO)
        
        # Abrir workbook
        wb = load_workbook(ruta)
        ws = wb.active
        if ws is None:
            logger.error("No se pudo abrir la hoja activa del workbook")
            raise ValueError("Hoja activa no encontrada")
        # Rellenar celdas básicas según mapeo
        campos_rellenados = 0
        for campo, celda in mapeo.items():
            # Saltar campo 'actividades' (se maneja aparte)
            if campo == 'actividades':
                continue
                
            if campo in datos and datos[campo]:
                try:
                    ws[celda] = datos[campo]
                    campos_rellenados += 1
                except Exception as e:
                    logger.warning(f"No se pudo rellenar celda {celda}: {e}")
        
        # Rellenar matriz de actividades
        if 'actividades' in datos and datos['actividades']:
            logger.info(f"Rellenando {len(datos['actividades'])} actividades")
            # TODO: Implementar lógica de tabla dinámica
            # Por ahora solo log
        
        logger.info(f"Campos rellenados: {campos_rellenados}/{len(mapeo)}")
        
        # Guardar y cerrar
        wb.save(ruta)
        wb.close()
        logger.info(f"✅ Instructivo guardado exitosamente")
        
    except Exception as e:
        logger.error(f"Error al rellenar instructivo: {e}")
        raise
