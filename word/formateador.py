from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from config import Config
from utils.logger import configurar_logger
from datetime import datetime
import copy

logger = configurar_logger("Word-Formateador")

def llenar_procedimiento(ruta: Path, datos: dict):
    """
    Rellena el formato Word de PROCEDIMIENTO (tablas estructuradas)
    """
    print(f"\n{'='*60}")
    print(f"[FORMATEADOR] Iniciando llenado de procedimiento")
    print(f"[FORMATEADOR] Archivo: {ruta.name}")
    print(f"[FORMATEADOR] Código: {datos.get('codigo', 'N/A')}")
    print(f"{'='*60}\n")
    
    logger.info(f"🔄 [WORD] Rellenando procedimiento: {ruta.name}")
    logger.info(f"   Código: {datos.get('codigo', 'N/A')}")
    logger.info(f"   Versión: {datos.get('version', 'N/A')}")
    
    try:
        doc = Document(str(ruta))
        logger.info(f"✅ Documento Word abierto correctamente")
        
        # ========== ENCABEZADO DE PÁGINA (HEADER WORD) ==========
        # La tabla del header tiene 2 filas x 2 columnas:
        #   [0, 0] = Logos (izquierda, no tocar)
        #   [0, 1] = TÍTULO del formato (nombre del proceso)
        #   [1, 0] = vacío (izquierda, no tocar)
        #   [1, 1] = ÁREA que realiza el formato
        try:
            section = doc.sections[0]
            header = section.header
            if header.tables:
                tabla_header = header.tables[0]
                # Fila 0, columna derecha → Título (nombre del proceso): negrita + centrado
                p_titulo = tabla_header.rows[0].cells[1].paragraphs[0]
                p_titulo.clear()
                run_titulo = p_titulo.add_run(datos.get('nombre_proceso', '').upper())
                run_titulo.bold = True
                p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Fila 1, columna derecha → Área que realiza: negrita + centrado
                p_area = tabla_header.rows[1].cells[1].paragraphs[0]
                p_area.clear()
                run_area = p_area.add_run(datos.get('area_realiza', '').upper())
                run_area.bold = True
                p_area.alignment = WD_ALIGN_PARAGRAPH.CENTER
                logger.info("✅ [Header] Encabezado Word rellenado (bold+center) correctamente")

                # ========== LOGOS: ACTUALIZAR CÓDIGO Y FECHA EN TEXTBOXES ==========
                # La celda izquierda del header (tc[0]) contiene 3 textboxes WPS + 3 VML (espejo)
                # p[0] = código del documento  (ej: PROC-002-LOG)
                # p[1] = fecha DD/MM/YYYY Vx   (ej: 19/02/2026 V1)
                try:
                    NS_WPS = 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'
                    NS_V   = 'urn:schemas-microsoft-com:vml'
                    W_NS   = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'

                    # Convertir fecha YYYY-MM-DD → DD/MM/YYYY
                    fecha_raw = datos.get('fecha', '')
                    try:
                        fecha_col = datetime.strptime(fecha_raw, '%Y-%m-%d').strftime('%d/%m/%Y')
                    except Exception:
                        fecha_col = fecha_raw  # usar tal cual si el formato es otro

                    version_str = str(datos.get('version', '1'))
                    if not version_str.upper().startswith('V'):
                        version_str = f"V{version_str}"

                    texto_p0 = datos.get('codigo', '')          # línea 1: código
                    texto_p1 = f"{fecha_col} {version_str}"     # línea 2: DD/MM/YYYY Vx

                    def _reemplazar_parrafo_txbx(p_elem, texto_nuevo):
                        """
                        Sustituye el contenido de un w:p dentro de un textbox.
                        Conserva el w:rPr del primer w:r existente, elimina todos
                        los w:r y crea uno nuevo con el texto deseado.
                        Aplica 7pt (val=14) para que el texto quepa sin desbordarse.
                        """
                        runs_ex = p_elem.findall(f'{{{W_NS}}}r')
                        rPr_copy = None
                        for r in runs_ex:
                            rPr = r.find(f'{{{W_NS}}}rPr')
                            if rPr is not None:
                                rPr_copy = copy.deepcopy(rPr)
                                break
                        # Ajustar tamaño de fuente a 7pt (val=14) para evitar desborde
                        if rPr_copy is not None:
                            for sz_tag in [qn('w:sz'), qn('w:szCs')]:
                                sz_elem = rPr_copy.find(sz_tag)
                                if sz_elem is not None:
                                    sz_elem.set(qn('w:val'), '14')
                                else:
                                    new_sz = OxmlElement(sz_tag.split('}')[1])
                                    new_sz.set(qn('w:val'), '14')
                                    rPr_copy.append(new_sz)
                        for r in runs_ex:
                            p_elem.remove(r)
                        nuevo_r = OxmlElement('w:r')
                        if rPr_copy is not None:
                            nuevo_r.append(rPr_copy)
                        nuevo_t = OxmlElement('w:t')
                        nuevo_t.text = texto_nuevo
                        if texto_nuevo and (texto_nuevo[0] == ' ' or texto_nuevo[-1] == ' '):
                            nuevo_t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                        nuevo_r.append(nuevo_t)
                        p_elem.append(nuevo_r)

                    # Celda izquierda del header (primer w:tc del primer w:tr)
                    tc0 = tabla_header.rows[0]._tr.findall(qn('w:tc'))[0]

                    # --- WPS textboxes (formato moderno) ---
                    wps_txbx_list = tc0.findall(f'.//{{{NS_WPS}}}txbx')
                    logger.info(f"[Logos-WPS] Textboxes encontrados: {len(wps_txbx_list)}")
                    for idx_tb, txbx in enumerate(wps_txbx_list):
                        parrafos = txbx.findall(f'.//{{{W_NS}}}p')
                        logger.info(f"  txbx_wps[{idx_tb}]: {len(parrafos)} párrafos")
                        if len(parrafos) >= 1:
                            _reemplazar_parrafo_txbx(parrafos[0], texto_p0)
                        if len(parrafos) >= 2:
                            _reemplazar_parrafo_txbx(parrafos[1], texto_p1)

                    # --- VML textboxes (espejo legacy) ---
                    vml_txbx_list = tc0.findall(f'.//{{{NS_V}}}textbox')
                    logger.info(f"[Logos-VML] Textboxes encontrados: {len(vml_txbx_list)}")
                    for idx_tb, txbx in enumerate(vml_txbx_list):
                        parrafos = txbx.findall(f'.//{{{W_NS}}}p')
                        logger.info(f"  txbx_vml[{idx_tb}]: {len(parrafos)} párrafos")
                        if len(parrafos) >= 1:
                            _reemplazar_parrafo_txbx(parrafos[0], texto_p0)
                        if len(parrafos) >= 2:
                            _reemplazar_parrafo_txbx(parrafos[1], texto_p1)

                    logger.info(f"✅ [Logos] Actualizados → código='{texto_p0}', fecha-ver='{texto_p1}'")
                    logger.info(f"   WPS: {len(wps_txbx_list)} textboxes | VML: {len(vml_txbx_list)} textboxes")
                except Exception as e_logos:
                    logger.warning(f"⚠️ [Logos] No se pudo actualizar textboxes: {e_logos}")
            else:
                logger.warning("⚠️ [Header] No se encontró tabla en el encabezado Word")
        except Exception as e_header:
            logger.warning(f"⚠️ [Header] No se pudo rellenar encabezado: {e_header}")
        
        # ========== TABLA 0: INFORMACIÓN BÁSICA ==========
        tabla_basica = doc.tables[0]
        logger.info(f"[Tabla 0] Filas: {len(tabla_basica.rows)}")
        
        # Rellenar AMBAS columnas (izquierda: título con numPr intacto, derecha: dato)
        # IMPORTANTE: para la columna izquierda usar p.clear()+add_run (preserva numPr)
        # Word numera automáticamente del 1 al 4 via la lista numérica de la plantilla
        if len(tabla_basica.rows) >= 4:
            # Fila 0
            tabla_basica.rows[0].cells[0].paragraphs[0].clear()
            tabla_basica.rows[0].cells[0].paragraphs[0].add_run("DENOMINACIÓN DEL PROCESO")
            tabla_basica.rows[0].cells[1].text = datos.get('nombre_proceso', '')
            
            # Fila 1
            tabla_basica.rows[1].cells[0].paragraphs[0].clear()
            tabla_basica.rows[1].cells[0].paragraphs[0].add_run("OBJETIVO")
            tabla_basica.rows[1].cells[1].text = datos.get('objetivo', '')
            
            # Fila 2
            tabla_basica.rows[2].cells[0].paragraphs[0].clear()
            tabla_basica.rows[2].cells[0].paragraphs[0].add_run("ALCANCE")
            tabla_basica.rows[2].cells[1].text = datos.get('alcance', '')
            
            # Fila 3
            tabla_basica.rows[3].cells[0].paragraphs[0].clear()
            tabla_basica.rows[3].cells[0].paragraphs[0].add_run("RESPONSABLES")
            tabla_basica.rows[3].cells[1].text = datos.get('responsables', '')
            
            logger.info("✅ [Tabla 0] Información básica rellenada")
        
        # ========== TABLA 1: TÉRMINOS Y CONDICIONES ==========
        # Estructura real de la plantilla (3 columnas en el XML):
        #   Fila 0: Header "5. TÉRMINOS Y DEFINICIONES" (tc[0] gridSpan=3, fusionada total)
        #   Fila 1: Plantilla término  → tc[0] gridSpan=2 (nombre) | tc[1] sin span (definición)
        #             En python-docx:    cells[0]=nombre,             cells[2]=definición
        #   Fila 2: Header "6. CONDICIONES GENERALES" (tc[0] gridSpan=3, fusionada total)
        #   Fila 3: Plantilla condición → tc[0] sin span (número) | tc[1] gridSpan=2 (descripción)
        #             En python-docx:    cells[0]=número,              cells[1]=descripción
        #
        # IMPORTANTE: add_row() rompe gridSpan → se debe clonar XML de la fila plantilla

        tabla_tc = doc.tables[1]
        tbl_elem = tabla_tc._tbl
        logger.info(f"[Tabla 1] Filas iniciales: {len(tabla_tc.rows)}")

        # Guardar XML de filas plantilla ANTES de cualquier modificación
        xml_fila_termino    = copy.deepcopy(tabla_tc.rows[1]._tr)
        xml_fila_condicion  = copy.deepcopy(tabla_tc.rows[3]._tr)

        # Referencia al tr del header de condiciones (para insertar términos antes de él)
        tr_header_condiciones = tabla_tc.rows[2]._tr

        def _escribir_celda(celda, texto):
            """Limpia la celda y escribe el texto preservando el formato de la plantilla."""
            for p in celda.paragraphs:
                p.clear()
            if not celda.paragraphs:
                celda._tc.append(OxmlElement('w:p'))
            celda.paragraphs[0].add_run(texto)

        def _escribir_en_tr(tr, idx_tc, texto):
            """Escribe texto en el tc número idx_tc del XML bruto de un w:tr clonado."""
            tcs = tr.findall(qn('w:tc'))
            if idx_tc >= len(tcs):
                return
            tc = tcs[idx_tc]
            # Limpiar todos los w:t existentes
            for t_elem in tc.findall('.//' + qn('w:t')):
                t_elem.text = ''
            # Escribir en el primer párrafo
            p_elems = tc.findall(qn('w:p'))
            if not p_elems:
                return
            p = p_elems[0]
            for r in p.findall(qn('w:r')):
                p.remove(r)
            r_new = OxmlElement('w:r')
            t_new = OxmlElement('w:t')
            t_new.text = texto
            if texto and (' ' in texto or texto != texto.strip()):
                t_new.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            r_new.append(t_new)
            p.append(r_new)

        # --- SECCIÓN: Términos y Definiciones ---
        terminos = datos.get('terminos', [])

        if terminos:
            # Fila 1 (plantilla): llenar con el primer término
            fila_t = tabla_tc.rows[1]
            _escribir_celda(fila_t.cells[0], terminos[0].get('nombre', ''))   # izquierda (gridSpan=2)
            _escribir_celda(fila_t.cells[2], terminos[0].get('definicion', ''))  # derecha

            # Términos adicionales: clonar XML plantilla e insertar antes del header de condiciones
            for td in terminos[1:]:
                nueva_tr = copy.deepcopy(xml_fila_termino)
                _escribir_en_tr(nueva_tr, 0, td.get('nombre', ''))      # tc[0] = nombre
                _escribir_en_tr(nueva_tr, 1, td.get('definicion', ''))  # tc[1] = definición
                tr_header_condiciones.addprevious(nueva_tr)

            logger.info(f"✅ [Tabla 1] Términos rellenados: {len(terminos)}")
        else:
            # Sin términos: dejar la fila plantilla vacía
            fila_t = tabla_tc.rows[1]
            _escribir_celda(fila_t.cells[0], '')
            _escribir_celda(fila_t.cells[2], '')

        # --- SECCIÓN: Condiciones Generales ---
        condiciones = datos.get('condiciones', [])

        if condiciones:
            # La fila plantilla de condición es ahora la última fila de la tabla
            ultima_fila = tabla_tc.rows[-1]
            _escribir_celda(ultima_fila.cells[0], '1')                                    # número
            _escribir_celda(ultima_fila.cells[1], condiciones[0].get('descripcion', ''))  # descripción (gridSpan=2)

            # Condiciones adicionales: clonar XML plantilla y agregar al final
            for idx, cd in enumerate(condiciones[1:], 2):
                nueva_tr = copy.deepcopy(xml_fila_condicion)
                _escribir_en_tr(nueva_tr, 0, str(idx))                     # tc[0] = número
                _escribir_en_tr(nueva_tr, 1, cd.get('descripcion', ''))    # tc[1] = descripción
                tbl_elem.append(nueva_tr)

            logger.info(f"✅ [Tabla 1] Condiciones rellenadas: {len(condiciones)}")
        else:
            # Sin condiciones: dejar la fila plantilla vacía
            ultima_fila = tabla_tc.rows[-1]
            _escribir_celda(ultima_fila.cells[0], '')
            _escribir_celda(ultima_fila.cells[1], '')

        logger.info(f"✅ [Tabla 1] Completa. Filas totales: {len(tabla_tc.rows)}")
        
        # ========== TABLA 2: FORMATOS Y NORMAS Y REQUISITOS LEGALES ==========
        if len(doc.tables) > 2:
            tabla_formatos = doc.tables[2]
            if len(tabla_formatos.rows) >= 2:
                # Header: p.clear()+add_run preserva numPr → Word lo numera como 7 automáticamente
                tabla_formatos.rows[0].cells[0].paragraphs[0].clear()
                tabla_formatos.rows[0].cells[0].paragraphs[0].add_run("FORMATOS, NORMAS Y REQUISITOS LEGALES")
                # Contenido: celda sin numPr, .text es seguro
                tabla_formatos.rows[1].cells[0].text = datos.get('formatos_normas', '')
                logger.info("✅ [Tabla 2] Formatos rellenados")
        
        # ========== TABLA 3: CONTROL DE REGISTROS ==========
        if len(doc.tables) > 3:
            tabla_control = doc.tables[3]
            if len(tabla_control.rows) >= 5:
                # Fila 2: ELABORÓ
                tabla_control.rows[2].cells[1].text = datos.get('elaboro_nombre', '')
                tabla_control.rows[2].cells[2].text = datos.get('elaboro_cargo', '')
                tabla_control.rows[2].cells[3].text = datos.get('elaboro_fecha', datos.get('fecha', ''))
                
                # Fila 3: REVISÓ
                tabla_control.rows[3].cells[1].text = datos.get('reviso_nombre', '')
                tabla_control.rows[3].cells[2].text = datos.get('reviso_cargo', '')
                tabla_control.rows[3].cells[3].text = datos.get('reviso_fecha', datos.get('fecha', ''))
                
                # Fila 4: APROBÓ
                tabla_control.rows[4].cells[1].text = datos.get('aprobo_nombre', '')
                tabla_control.rows[4].cells[2].text = datos.get('aprobo_cargo', '')
                tabla_control.rows[4].cells[3].text = datos.get('aprobo_fecha', datos.get('fecha', ''))
                
                logger.info("✅ [Tabla 3] Control de registros rellenado")
        
        # ========== TABLA 4: CONTROL DE CAMBIOS ==========
        if len(doc.tables) > 4:
            tabla_cambios = doc.tables[4]
            if len(tabla_cambios.rows) >= 3:
                tabla_cambios.rows[2].cells[0].text = "1"
                tabla_cambios.rows[2].cells[1].text = datos.get('fecha', datetime.now().strftime(Config.FECHA_FORMATO))
                tabla_cambios.rows[2].cells[2].text = "Versión inicial"
                logger.info("✅ [Tabla 4] Control de cambios rellenado")
        
        # Guardar documento
        doc.save(ruta)
        logger.info(f"💾 [WORD] Guardando documento...")
        logger.info(f"✅ [WORD] Procedimiento guardado exitosamente: {ruta.name}")
        
        print(f"\n✅ DOCUMENTO GENERADO: {ruta.name}\n")
        
    except Exception as e:
        error_msg = f"❌ ERROR CRÍTICO en formateador: {str(e)}"
        print(f"\n{error_msg}\n")
        logger.error(error_msg)
        import traceback
        tb = traceback.format_exc()
        logger.error(f"Traceback completo:\n{tb}")
        print(f"Traceback:\n{tb}")
        raise

def llenar_instructivo(ruta: Path, datos: dict):
    """
    Rellena el formato Word de INSTRUCTIVO (igual que procedimiento)
    """
    # Usar misma lógica que procedimiento
    return llenar_procedimiento(ruta, datos)
